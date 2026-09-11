"""Victus AI - Excel Automation"""
import os

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))

HAS_EXCEL = False
try:
    from openpyxl import Workbook, load_workbook
    HAS_EXCEL = True
except ImportError:
    pass


def excel_path():
    from config import EXCEL_FILE
    return os.path.join(SCRIPT_DIR, EXCEL_FILE)


def handle_excel(text):
    """Handle Excel commands. Returns (handled, response)."""
    if not HAS_EXCEL:
        return False, "Excel module not installed. Run pip install openpyxl"

    if "create excel" in text or "new excel" in text or "make excel" in text:
        wb = Workbook()
        wb.save(excel_path())
        return True, "Excel file created"

    if "write" in text and "excel" in text:
        # Extract what to write (everything after 'write')
        content = text.split("write", 1)[-1].replace("in excel", "").replace("to excel", "").strip()
        if not content:
            return True, "What should I write?"

        path = excel_path()
        if not os.path.exists(path):
            wb = Workbook()
            wb.save(path)

        wb = load_workbook(path)
        ws = wb.active
        # Find next empty row
        row = ws.max_row + 1 if ws.max_row and ws.cell(1, 1).value else 1
        ws.cell(row=row, column=1, value=content)
        wb.save(path)
        return True, f"Written to Excel row {row}"

    if "read excel" in text or "show excel" in text:
        path = excel_path()
        if not os.path.exists(path):
            return True, "No Excel file found. Say create excel first."
        wb = load_workbook(path)
        ws = wb.active
        rows = []
        for row in ws.iter_rows(max_row=5, values_only=True):
            rows.append(", ".join(str(c) for c in row if c))
        if rows:
            return True, "Excel data: " + ". ".join(rows)
        return True, "Excel is empty"

    return False, ""
